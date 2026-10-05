"""Excel import (Feature 3): read the workbook, validate every row, then upsert by SKU.

The files come from the customer and contain mistakes on purpose, so each row is
checked on its own: valid rows are imported, invalid rows are reported back with
their Excel row number and the reason, and never stop the rest of the file.
"""

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

from fastapi import HTTPException, status
from openpyxl import load_workbook

from app.repositories.product_repository import ProductRepository
from app.schemas.product import Category, FieldChange, ImportIssue, ImportResult, ImportRow

CATEGORIES: tuple[str, ...] = Category.__args__
STATUSES = ("active", "inactive")
SKU_PATTERN = re.compile(r"^LS-\d{4}$")

COLUMNS = ("sku", "name", "category", "price", "size", "description", "how_to_use", "status")
REQUIRED_COLUMNS = ("sku", "name", "category", "price")
# Fields compared with the database to tell "update" from "unchanged"
COMPARED_FIELDS = COLUMNS[1:]

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 5000
MAX_PRICE = Decimal("99999999.99")  # numeric(10, 2)
MAX_NAME_LENGTH = 200
MAX_TEXT_LENGTH = 5000


@dataclass
class ParsedRow:
    row: int
    sku: str | None = None
    data: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# --- cell helpers ----------------------------------------------------------------


def _text(value) -> str | None:
    """Cell -> trimmed string, or None when empty. 30.0 -> "30"."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _parse_sku(value, row: ParsedRow) -> None:
    raw = _text(value)
    if raw is None:
        row.errors.append("ไม่มี SKU")
        return
    sku = raw.upper()
    if not SKU_PATTERN.match(sku):
        row.errors.append(f'SKU "{raw}" ไม่ตรงรูปแบบ LS-0000 (LS- ตามด้วยตัวเลข 4 หลัก)')
        return
    if sku != str(value):
        row.warnings.append(f'ปรับ SKU "{value}" เป็น "{sku}" (ตัดช่องว่าง/เปลี่ยนเป็นตัวพิมพ์ใหญ่)')
    row.sku = sku


def _parse_name(value, row: ParsedRow) -> None:
    name = _text(value)
    if name is None:
        row.errors.append("ไม่มีชื่อสินค้า")
    elif len(name) > MAX_NAME_LENGTH:
        row.errors.append(f"ชื่อสินค้ายาวเกิน {MAX_NAME_LENGTH} ตัวอักษร")
    else:
        row.data["name"] = name


def _parse_category(value, row: ParsedRow) -> None:
    raw = _text(value)
    if raw is None:
        row.errors.append("ไม่มีหมวดหมู่ (category)")
        return
    match = next((c for c in CATEGORIES if c.lower() == raw.lower()), None)
    if match is None:
        row.errors.append(f'หมวดหมู่ "{raw}" ไม่ถูกต้อง ต้องเป็นหนึ่งใน {", ".join(CATEGORIES)}')
        return
    row.data["category"] = match


def _parse_price(value, row: ParsedRow) -> None:
    if value is None or (isinstance(value, str) and not value.strip()):
        row.errors.append("ไม่มีราคา")
        return
    if isinstance(value, bool):
        row.errors.append(f'ราคา "{value}" ไม่ใช่ตัวเลข')
        return

    if isinstance(value, (int, float)):
        cleaned = str(value)
    else:
        # Common ways people type prices: "1,290", "฿590", "590 บาท"
        cleaned = re.sub(r"[,\s฿]|บาท", "", str(value))
    try:
        price = Decimal(cleaned)
    except InvalidOperation:
        row.errors.append(f'ราคา "{value}" ไม่ใช่ตัวเลข')
        return
    if not price.is_finite() or price <= 0:
        row.errors.append(f"ราคาต้องมากกว่า 0 (พบ {value})")
        return
    if price > MAX_PRICE:
        row.errors.append(f"ราคาสูงเกินไป (พบ {value})")
        return

    price = price.quantize(Decimal("0.01"))
    if isinstance(value, str):
        row.warnings.append(f'แปลงราคา "{value}" เป็น {price.normalize():f}')
    row.data["price"] = float(price)


def _parse_status(value, row: ParsedRow) -> None:
    raw = _text(value)
    if raw is None:
        row.data["status"] = "active"  # requirement: empty status = active
        return
    if raw.lower() not in STATUSES:
        row.errors.append(f'สถานะ "{raw}" ไม่ถูกต้อง ต้องเป็น active หรือ inactive (เว้นว่าง = active)')
        return
    row.data["status"] = raw.lower()


def _parse_optional_text(column: str, value, row: ParsedRow) -> None:
    text = _text(value)
    if text is not None and len(text) > MAX_TEXT_LENGTH:
        row.errors.append(f"{column} ยาวเกิน {MAX_TEXT_LENGTH} ตัวอักษร")
        return
    row.data[column] = text


# --- workbook ----------------------------------------------------------------------


def _bad_file(message: str) -> HTTPException:
    return HTTPException(status.HTTP_400_BAD_REQUEST, message)


def _read_rows(content: bytes) -> list[tuple]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as err:
        raise _bad_file("ไม่สามารถเปิดไฟล์ได้ กรุณาใช้ไฟล์ Excel (.xlsx) ที่ไม่ได้ตั้งรหัสผ่าน") from err
    try:
        sheet = workbook.worksheets[0]
        return [tuple(r) for r in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _column_index(header: tuple) -> dict[str, int]:
    """Maps column name -> position. Headers are matched case/space-insensitively."""
    index: dict[str, int] = {}
    for position, cell in enumerate(header):
        name = (_text(cell) or "").lower().replace(" ", "_")
        if name in COLUMNS and name not in index:
            index[name] = position
    missing = [c for c in REQUIRED_COLUMNS if c not in index]
    if missing:
        raise _bad_file(
            f"ไม่พบคอลัมน์ {', '.join(missing)} ในแถวแรกของชีต "
            f"(แถวแรกต้องเป็นหัวตาราง: {', '.join(COLUMNS)})"
        )
    return index


def parse_workbook(content: bytes) -> list[ParsedRow]:
    """Validates every non-empty data row. Pure: does not touch the database."""
    rows = _read_rows(content)
    if not rows:
        raise _bad_file("ไฟล์ไม่มีข้อมูล")
    index = _column_index(rows[0])

    parsed: list[ParsedRow] = []
    first_row_of_sku: dict[str, int] = {}
    for excel_row, cells in enumerate(rows[1:], start=2):
        if all(_text(c) is None for c in cells):
            continue  # blank line in the sheet
        if len(parsed) >= MAX_DATA_ROWS:
            raise _bad_file(f"ไฟล์มีข้อมูลเกิน {MAX_DATA_ROWS} แถว กรุณาแบ่งเป็นหลายไฟล์")

        def cell(column: str):
            position = index.get(column)
            return cells[position] if position is not None and position < len(cells) else None

        row = ParsedRow(row=excel_row)
        _parse_sku(cell("sku"), row)
        _parse_name(cell("name"), row)
        _parse_category(cell("category"), row)
        _parse_price(cell("price"), row)
        for column in ("size", "description", "how_to_use"):
            _parse_optional_text(column, cell(column), row)
        _parse_status(cell("status"), row)

        # The same SKU twice in one file is ambiguous: keep the first row, reject the rest
        if row.sku:
            first = first_row_of_sku.setdefault(row.sku, excel_row)
            if first != excel_row:
                row.errors.append(f"SKU {row.sku} ซ้ำกับแถวที่ {first} (ใช้ข้อมูลจากแถวที่ {first})")

        parsed.append(row)
    return parsed


# --- import ------------------------------------------------------------------------


def _same(field_name: str, old, new) -> bool:
    if field_name == "price":
        return old is not None and round(float(old), 2) == round(float(new), 2)
    return (old or None) == (new or None)


def _diff(existing: dict, new: dict) -> list[FieldChange]:
    return [
        FieldChange(field=f, old=existing.get(f), new=new.get(f))
        for f in COMPARED_FIELDS
        if not _same(f, existing.get(f), new.get(f))
    ]


class ProductImportService:
    def __init__(self, repo: ProductRepository):
        self.repo = repo

    def run(self, file_name: str, content: bytes, dry_run: bool) -> ImportResult:
        if not file_name.lower().endswith(".xlsx"):
            raise _bad_file("รองรับเฉพาะไฟล์ .xlsx")
        if not content:
            raise _bad_file("ไฟล์ว่างเปล่า")
        if len(content) > MAX_FILE_BYTES:
            raise _bad_file(f"ไฟล์ใหญ่เกิน {MAX_FILE_BYTES // (1024 * 1024)} MB")

        parsed = parse_workbook(content)
        valid = [r for r in parsed if not r.errors]
        existing = self.repo.get_import_fields_by_skus([r.sku for r in valid])

        rows: list[ImportRow] = []
        to_write: list[dict] = []
        for r in valid:
            payload = {"sku": r.sku, **r.data}
            current = existing.get(r.sku)
            changes = _diff(current, payload) if current else []
            action = "create" if current is None else ("update" if changes else "unchanged")
            rows.append(ImportRow(row=r.row, sku=r.sku, name=r.data["name"], action=action, changes=changes))
            if action != "unchanged":
                to_write.append(payload)

        if not dry_run and to_write:
            self.repo.upsert_many(to_write)

        return ImportResult(
            dry_run=dry_run,
            file_name=file_name,
            total_rows=len(parsed),
            created=sum(r.action == "create" for r in rows),
            updated=sum(r.action == "update" for r in rows),
            unchanged=sum(r.action == "unchanged" for r in rows),
            failed=len(parsed) - len(valid),
            rows=rows,
            errors=[ImportIssue(row=r.row, sku=r.sku, messages=r.errors) for r in parsed if r.errors],
            warnings=[
                ImportIssue(row=r.row, sku=r.sku, messages=r.warnings)
                for r in parsed
                if r.warnings and not r.errors
            ],
        )
