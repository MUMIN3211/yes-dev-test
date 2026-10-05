"""In-memory stand-ins for Supabase, used by the tests.

FakeSupabase implements the small part of the supabase-py query builder that
the repositories use (select/insert/upsert/update + eq/in_/or_/order/limit), so
integration tests run the real routers, services and repositories without a
network or a real database. FakeAuthGateway replaces Supabase Auth.
"""

import copy
import uuid
from datetime import datetime, timezone
from types import SimpleNamespace

from supabase_auth.errors import AuthApiError


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _product_defaults(row: dict) -> dict:
    row.setdefault("id", str(uuid.uuid4()))
    row.setdefault("image_url", None)
    row.setdefault("created_at", _now())
    # Mirrors the normalize_product_status trigger in supabase/schema.sql
    row["status"] = (row.get("status") or "").strip().lower() or "active"
    return row


def _admin_defaults(row: dict) -> dict:
    for key, value in {
        "is_active": True,
        "invited_by": None,
        "invited_at": None,
        "activated_at": None,
        "last_login_at": None,
        "created_at": _now(),
    }.items():
        row.setdefault(key, value)
    return row


DEFAULTS = {"products": _product_defaults, "admins": _admin_defaults}


class _Query:
    def __init__(self, db: "FakeSupabase", table: str):
        self.db = db
        self.table = table
        self.mode = "select"
        self.columns: list[str] | None = None
        self.payload = None
        self.on_conflict: str | None = None
        self.filters: list = []
        self.order_by: str | None = None
        self.max_rows: int | None = None

    # --- verbs ---
    def select(self, columns: str = "*"):
        self.mode = "select"
        self.columns = None if columns.strip() == "*" else [c.strip() for c in columns.split(",")]
        return self

    def insert(self, row):
        self.mode, self.payload = "insert", row
        return self

    def upsert(self, rows, on_conflict: str = "id"):
        self.mode, self.payload, self.on_conflict = "upsert", rows, on_conflict
        return self

    def update(self, values: dict):
        self.mode, self.payload = "update", values
        return self

    # --- filters ---
    def eq(self, column, value):
        self.filters.append(lambda r: r.get(column) == value)
        return self

    def in_(self, column, values):
        values = set(values)
        self.filters.append(lambda r: r.get(column) in values)
        return self

    def or_(self, expression: str):
        # Only the `col.ilike.%term%` form used by ProductRepository.list_active
        conditions = []
        for part in expression.split(","):
            column, op, pattern = part.split(".", 2)
            assert op == "ilike", op
            needle = pattern.strip("%").lower()
            conditions.append((column, needle))
        self.filters.append(lambda r: any(needle in str(r.get(col) or "").lower() for col, needle in conditions))
        return self

    def order(self, column: str, desc: bool = False):
        self.order_by = column
        return self

    def limit(self, n: int):
        self.max_rows = n
        return self

    # --- run ---
    def _matches(self, row) -> bool:
        return all(f(row) for f in self.filters)

    def _project(self, row: dict) -> dict:
        row = copy.deepcopy(row)
        return row if self.columns is None else {c: row.get(c) for c in self.columns}

    def execute(self):
        rows = self.db.tables.setdefault(self.table, [])
        self.db.calls.append((self.table, self.mode))
        defaults = DEFAULTS.get(self.table, lambda r: r)

        if self.mode == "select":
            result = [r for r in rows if self._matches(r)]
            if self.order_by:
                result.sort(key=lambda r: (r.get(self.order_by) is None, r.get(self.order_by)))
            if self.max_rows is not None:
                result = result[: self.max_rows]
            return SimpleNamespace(data=[self._project(r) for r in result])

        if self.mode == "insert":
            new = [defaults(dict(r)) for r in (self.payload if isinstance(self.payload, list) else [self.payload])]
            for row in new:
                self.db.check_unique(self.table, row)
            rows.extend(new)
            return SimpleNamespace(data=copy.deepcopy(new))

        if self.mode == "upsert":
            written = []
            for incoming in self.payload if isinstance(self.payload, list) else [self.payload]:
                key = incoming[self.on_conflict]
                existing = next((r for r in rows if r.get(self.on_conflict) == key), None)
                if existing is None:
                    row = defaults(dict(incoming))
                    rows.append(row)
                else:
                    existing.update(incoming)  # only the columns sent are changed
                    defaults(existing)
                    existing["updated_at"] = _now()
                    row = existing
                written.append(copy.deepcopy(row))
            return SimpleNamespace(data=written)

        if self.mode == "update":
            updated = []
            for row in rows:
                if self._matches(row):
                    row.update(self.payload)
                    defaults(row)
                    row["updated_at"] = _now()
                    updated.append(copy.deepcopy(row))
            return SimpleNamespace(data=updated)

        raise AssertionError(self.mode)


class FakeSupabase:
    UNIQUE = {"products": ("sku",), "admins": ("id", "email")}

    def __init__(self):
        self.tables: dict[str, list[dict]] = {"products": [], "admins": []}
        self.calls: list[tuple[str, str]] = []

    def table(self, name: str) -> _Query:
        return _Query(self, name)

    def check_unique(self, table: str, row: dict) -> None:
        for column in self.UNIQUE.get(table, ()):
            if any(r.get(column) == row.get(column) for r in self.tables[table]):
                raise ValueError(f"duplicate key {table}.{column}={row.get(column)}")

    # --- helpers for arranging test data ---
    def add_product(self, **fields) -> dict:
        row = {
            "sku": "LS-0001",
            "name": "Test Product",
            "category": "Serum",
            "price": 500,
            "size": "30 ml",
            "description": "desc",
            "how_to_use": "use",
            "status": "active",
            "updated_at": _now(),
            **fields,
        }
        self.tables["products"].append(_product_defaults(row))
        return row

    def writes(self, table: str) -> list[str]:
        return [mode for t, mode in self.calls if t == table and mode != "select"]


class FakeAuthGateway:
    """Supabase Auth (auth.users) in memory. Deleting a user also deletes its
    admins row, like the `on delete cascade` foreign key."""

    def __init__(self, db: FakeSupabase):
        self.db = db
        self.users: dict[str, dict] = {}  # id -> {"email", "password"}
        self.tokens: dict[str, str] = {}  # invitation access token -> user id
        self.invites_sent: list[tuple[str, str]] = []
        self.fail_invite_with: str | None = None

    def add_user(self, email: str, password: str | None = None) -> str:
        user_id = str(uuid.uuid4())
        self.users[user_id] = {"email": email, "password": password}
        return user_id

    def verify_password(self, email: str, password: str) -> str | None:
        for user_id, user in self.users.items():
            if user["email"] == email and user["password"] is not None and user["password"] == password:
                return user_id
        return None

    def invite(self, email: str, redirect_to: str) -> str:
        if self.fail_invite_with:
            raise AuthApiError(self.fail_invite_with, 400, None)
        user_id = self.add_user(email)
        self.tokens[f"invite-token-{user_id}"] = user_id
        self.invites_sent.append((email, redirect_to))
        return user_id

    def user_from_access_token(self, access_token: str):
        user_id = self.tokens.get(access_token)
        if user_id is None or user_id not in self.users:
            return None
        return SimpleNamespace(id=user_id, email=self.users[user_id]["email"])

    def set_password(self, user_id: str, password: str) -> None:
        self.users[user_id]["password"] = password

    def create_confirmed_user(self, email: str, password: str) -> str:
        return self.add_user(email, password)

    def find_user_id_by_email(self, email: str) -> str | None:
        return next((uid for uid, u in self.users.items() if u["email"] == email), None)

    def delete_user(self, user_id: str) -> None:
        self.users.pop(user_id, None)
        self.db.tables["admins"] = [r for r in self.db.tables["admins"] if r["id"] != user_id]

    def token_for(self, user_id: str) -> str:
        return next(t for t, uid in self.tokens.items() if uid == user_id)


def make_xlsx(rows: list[list], header: list[str] | None = None) -> bytes:
    """Builds an .xlsx in memory. header=None uses the standard 8 columns."""
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        header
        if header is not None
        else ["sku", "name", "category", "price", "size", "description", "how_to_use", "status"]
    )
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def valid_row(sku="LS-2001", name="Serum A", category="Serum", price=500, status="active", **extra) -> list:
    return [
        sku,
        name,
        category,
        price,
        extra.get("size", "30 ml"),
        extra.get("description", "desc"),
        extra.get("how_to_use", "use"),
        status,
    ]

